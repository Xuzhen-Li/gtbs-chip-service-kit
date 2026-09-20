"""OIV phenotype ETL: parse ST1 xlsx → long phenotype.tsv.

Rules (plan §2.2): 0/empty → missing; multi-value ``a|b`` ordinal mean / nominal
consensus; out-of-scale dropped; des-cep replicates median/mode; sources not merged.
"""

from __future__ import annotations

import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd

MIN_PHENO_N = 50

SOURCE_SHEETS = {
    "www.eu-vitis.de": "euvitis",
    "des-cep-monde-edition-2009": "descep",
}

BINARY_RULES: dict[str, str] = {
    "OIV 225": "1->0; 2,3,4,5,6,7->1",
    "OIV 236": "1->0; 2->1; else->NA",
    "OIV 151": "3->0; 4,5->1; 1,2->NA",
    "OIV 241": "1,2->1; 3->0",
}

# Compiled from BINARY_RULES strings.
_BINARY_MAP: dict[str, dict[float, float | None]] = {
    "OIV 225": {1.0: 0.0, 2.0: 1.0, 3.0: 1.0, 4.0: 1.0, 5.0: 1.0, 6.0: 1.0, 7.0: 1.0},
    "OIV 236": {1.0: 0.0, 2.0: 1.0},  # else NA
    "OIV 151": {3.0: 0.0, 4.0: 1.0, 5.0: 1.0},  # 1,2 NA
    "OIV 241": {1.0: 1.0, 2.0: 1.0, 3.0: 0.0},
}

_STRICT_BINARY_ELSE_NA = {"OIV 236", "OIV 151"}


def norm_header(c: Any) -> str:
    if c is None or (isinstance(c, float) and np.isnan(c)):
        return ""
    return str(c).replace("\n", " ").replace("\r", " ").strip()


def load_panel_ids(info_path: Path) -> set[str]:
    ids: set[str] = set()
    with info_path.open(encoding="utf-8") as fh:
        next(fh, None)
        for line in fh:
            if not line.strip():
                continue
            ids.add(line.split("\t", 1)[0].strip())
    return ids


def parse_allowed_from_notations(notations: str) -> set[int]:
    """Integers and a-b ranges from OIV notation text (e.g. ``1-3=...``)."""
    text = notations or ""
    allowed: set[int] = set()
    for m in re.finditer(r"(\d+)\s*-\s*(\d+)", text):
        a, b = int(m.group(1)), int(m.group(2))
        if a > b:
            a, b = b, a
        if b - a <= 20:
            allowed.update(range(a, b + 1))
    for m in re.finditer(r"(?<![\d.])(\d+)(?![\d.])", text):
        allowed.add(int(m.group(1)))
    return allowed or set(range(1, 11))


def parse_oiv_value(raw: Any, scale: str, allowed: set[int]) -> float | None:
    if raw is None:
        return None
    if isinstance(raw, float) and np.isnan(raw):
        return None
    s = str(raw).strip()
    if s == "" or s == "0" or s.lower() == "none" or s.lower() == "nan":
        return None
    parts = [p.strip() for p in s.replace(";", "|").split("|") if p.strip()]
    vals: list[float] = []
    for p in parts:
        try:
            v = float(p)
        except ValueError:
            continue
        if abs(v - round(v)) > 1e-9:
            continue
        iv = int(round(v))
        if iv not in allowed:
            continue
        vals.append(float(iv))
    if not vals:
        return None
    if len(vals) == 1:
        return vals[0]
    if scale == "ordinal":
        return float(sum(vals) / len(vals))
    if all(v == vals[0] for v in vals):
        return vals[0]
    return None


def aggregate_replicates(values: Iterable[float | None], scale: str) -> float | None:
    vals = [float(v) for v in values if v is not None]
    if not vals:
        return None
    if scale == "ordinal":
        return float(np.median(vals))
    counts = Counter(vals)
    most = counts.most_common()
    if len(most) == 1 or most[0][1] > most[1][1]:
        return float(most[0][0])
    return None


def apply_binary_rule(value: float | None, trait: str) -> float | None:
    if value is None:
        return None
    mapping = _BINARY_MAP.get(trait)
    if mapping is None:
        return None
    if value in mapping:
        return mapping[value]
    if trait in _STRICT_BINARY_ELSE_NA:
        return None
    return mapping.get(value)


def infer_scale(trait_type: str, notations: str) -> str:
    tt = (trait_type or "").strip().lower()
    if tt.startswith("quant"):
        return "ordinal"
    if tt.startswith("qual"):
        return "nominal"
    allowed = parse_allowed_from_notations(notations)
    ladder = {1, 3, 5, 7, 9}
    if ladder.issubset(allowed) or allowed == ladder:
        return "ordinal"
    return "nominal"


def trait_slug(trait: str, binary: bool = False) -> str:
    s = trait.replace(" ", "_")
    if binary:
        s += "_bin"
    return s


def find_col(headers: list[str], *needles: str) -> int | None:
    low = [h.lower() for h in headers]
    for i, h in enumerate(low):
        if all(n.lower() in h for n in needles):
            return i
    return None


def load_oiv_dictionary(euvitis_xlsx: Path, descep_xlsx: Path) -> pd.DataFrame:
    rows: list[dict[str, str]] = []
    seen: set[str] = set()

    def _read_dict(path: Path, source_sheet: str, has_type: bool) -> None:
        xl = pd.ExcelFile(path, engine="openpyxl")
        name = "Descriptor of OIV trait"
        if name not in xl.sheet_names:
            name = xl.sheet_names[0]
        df = pd.read_excel(xl, sheet_name=name, dtype=str)
        df.columns = [norm_header(c) for c in df.columns]
        cmap = {c.lower(): c for c in df.columns}

        def col(*keys: str) -> str | None:
            for k in keys:
                if k.lower() in cmap:
                    return cmap[k.lower()]
            for ck, orig in cmap.items():
                if all(x.lower() in ck for x in k.split()):
                    return orig
            return None

        c_code = col("OIV Code") or col("OIV")
        c_desc = col("Descriptor")
        c_not = col("Descriptor notations") or col("notations")
        c_type = col("Trait type") if has_type else None
        if c_code is None:
            return
        for _, r in df.iterrows():
            code = norm_header(r.get(c_code))
            if not code.startswith("OIV"):
                continue
            desc = norm_header(r.get(c_desc) if c_desc else "")
            notations = str(r.get(c_not) if c_not else "") if c_not else ""
            ttype = norm_header(r.get(c_type) if c_type else "")
            scale = infer_scale(ttype, notations)
            note = notations
            if source_sheet == "des-cep-monde-edition-2009" and code in seen:
                continue
            if source_sheet == "des-cep-monde-edition-2009" and code not in seen:
                note = (notations + " [待确认] des-cep-only dictionary").strip()
                ttype = ttype or "unknown"
            br = BINARY_RULES.get(code, "")
            rows.append(
                {
                    "code": code,
                    "descriptor": desc,
                    "notations": note.replace("\n", " / "),
                    "trait_type": ttype or ("Quantitative" if scale == "ordinal" else "Qualitative"),
                    "scale": scale,
                    "binary_rule": br,
                    "source_sheet": source_sheet,
                }
            )
            seen.add(code)

    if euvitis_xlsx.exists():
        _read_dict(euvitis_xlsx, "www.eu-vitis.de", has_type=True)
    if descep_xlsx.exists():
        _read_dict(descep_xlsx, "des-cep-monde-edition-2009", has_type=False)
    return pd.DataFrame(rows)


def _scale_for(code: str, dict_df: pd.DataFrame) -> str:
    hit = dict_df.loc[dict_df["code"] == code]
    if len(hit):
        return str(hit.iloc[0]["scale"])
    return "ordinal"


def _allowed_for(code: str, dict_df: pd.DataFrame) -> set[int]:
    hit = dict_df.loc[dict_df["code"] == code]
    if len(hit):
        return parse_allowed_from_notations(str(hit.iloc[0]["notations"]))
    return set(range(1, 11))


def sheet_excel_data_rows(path: Path, sheet: str) -> int:
    """Excel used-range data rows (max_row minus header), including blanks.

    pandas ``read_excel`` drops trailing empty rows, so this can be larger than
    ``len(df)``. ST1 des-cep: max_row−1 = 2256, nonempty Library-ID rows = 312.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        return max(int(ws.max_row or 1) - 1, 0)
    finally:
        wb.close()


def extract_st1_sheet(
    path: Path,
    sheet: str,
    panel_ids: set[str],
    dict_df: pd.DataFrame,
) -> tuple[list[dict], list[dict], dict[str, Any]]:
    """Return phenotype rows, annot rows, and per-sheet counts."""
    df = pd.read_excel(path, sheet_name=sheet, dtype=str, engine="openpyxl")
    df.columns = [norm_header(c) for c in df.columns]
    headers = list(df.columns)
    lib_i = find_col(headers, "library", "id")
    if lib_i is None:
        lib_i = 1 if len(headers) > 1 else 0
    lib_col = headers[lib_i]
    oiv_cols = [c for c in headers if c.startswith("OIV")]

    stats = {
        "xlsx_max_row": sheet_excel_data_rows(path, sheet),
        "xlsx_rows": int(len(df)),
        "uniq_libid": int(df[lib_col].dropna().astype(str).str.strip().nunique()),
        "cells_total": 0,
        "cells_zero_dropped": 0,
        "cells_multivalue": 0,
        "cells_out_of_scale_dropped": 0,
        "ids_in_panel": set(),
    }

    # phenotype accum: (id, trait) -> list of (raw, parsed)
    acc: dict[tuple[str, str], list[tuple[str, float | None]]] = defaultdict(list)
    annot_rows: list[dict] = []

    v_num = find_col(headers, "vivc", "number")
    v_prime = find_col(headers, "prime", "name") or find_col(headers, "vivc prime")
    v_uti = find_col(headers, "utilization")
    v_sex = find_col(headers, "flower", "sex")
    v_mus = find_col(headers, "muscat")
    v_skin = find_col(headers, "color", "berry") or find_col(headers, "berry skin")

    for _, row in df.iterrows():
        lib = norm_header(row.get(lib_col))
        if not lib:
            continue
        in_panel = lib in panel_ids
        if in_panel:
            stats["ids_in_panel"].add(lib)
        for c in oiv_cols:
            raw = row.get(c)
            raw_s = "" if raw is None or (isinstance(raw, float) and np.isnan(raw)) else str(raw).strip()
            if raw_s == "":
                continue
            stats["cells_total"] += 1
            if raw_s == "0":
                stats["cells_zero_dropped"] += 1
            if "|" in raw_s:
                stats["cells_multivalue"] += 1
            if not in_panel:
                continue
            scale = _scale_for(c, dict_df)
            allowed = _allowed_for(c, dict_df)
            parsed = parse_oiv_value(raw_s, scale, allowed)
            if parsed is None and raw_s not in ("0", ""):
                stats["cells_out_of_scale_dropped"] += 1
            acc[(lib, c)].append((raw_s, parsed))

        if in_panel:
            annot_rows.append(
                {
                    "ID": lib,
                    "vivc_number": norm_header(row.iloc[v_num]) if v_num is not None else "",
                    "vivc_prime_name": norm_header(row.iloc[v_prime]) if v_prime is not None else "",
                    "vivc_utilization": norm_header(row.iloc[v_uti]) if v_uti is not None else "",
                    "vivc_flower_sex": norm_header(row.iloc[v_sex]) if v_sex is not None else "",
                    "vivc_muscat": norm_header(row.iloc[v_mus]) if v_mus is not None else "",
                    "berry_skin_color_text": norm_header(row.iloc[v_skin]) if v_skin is not None else "",
                }
            )

    source = SOURCE_SHEETS[sheet]
    pheno_rows: list[dict] = []
    replicates_aggregated = 0
    for (sid, trait), pairs in acc.items():
        scale = _scale_for(trait, dict_df)
        parsed_vals = [p for _, p in pairs]
        if len(pairs) > 1:
            replicates_aggregated += 1
        value = aggregate_replicates(parsed_vals, scale)
        if value is None:
            continue
        raw_join = "|".join(r for r, _ in pairs)
        pheno_rows.append(
            {
                "ID": sid,
                "trait": trait,
                "value": value,
                "source": source,
                "n_obs": len(pairs),
                "raw": raw_join,
            }
        )
    stats["replicates_aggregated"] = replicates_aggregated
    stats["n_pheno_rows"] = len(pheno_rows)
    stats["ids_in_panel_n"] = len(stats["ids_in_panel"])
    return pheno_rows, annot_rows, stats


def merge_annot(euvitis_ann: list[dict], descep_ann: list[dict]) -> list[dict]:
    by_id: dict[str, dict] = {}
    for row in descep_ann:
        r = dict(row)
        r["source"] = "descep"
        by_id[r["ID"]] = r
    for row in euvitis_ann:
        r = dict(row)
        r["source"] = "euvitis"
        old = by_id.get(r["ID"])
        if old:
            for k in (
                "vivc_number",
                "vivc_prime_name",
                "vivc_utilization",
                "vivc_flower_sex",
                "vivc_muscat",
                "berry_skin_color_text",
            ):
                a, b = str(old.get(k, "")), str(r.get(k, ""))
                if a and b and a != b:
                    r[k] = f"{b} [conflict]"
                elif not b:
                    r[k] = a
        by_id[r["ID"]] = r
    return list(by_id.values())


def coverage_table(pheno: pd.DataFrame, dict_df: pd.DataFrame, panel_n: dict[str, int]) -> pd.DataFrame:
    rows = []
    scale_map = dict(zip(dict_df["code"], dict_df["scale"])) if len(dict_df) else {}
    for (source, trait), sub in pheno.groupby(["source", "trait"]):
        vals = sub["value"].astype(float)
        dist = Counter(vals.round(3).tolist())
        top = ",".join(f"{k}:{v}" for k, v in dist.most_common(10))
        rows.append(
            {
                "source": source,
                "trait": trait,
                "n_in_panel": panel_n.get(source, 0),
                "n_nonmissing": int(len(sub)),
                "scale": scale_map.get(trait, ""),
                "value_top10": top,
            }
        )
    return pd.DataFrame(rows).sort_values(["source", "n_nonmissing"], ascending=[True, False])


def load_phenotype_table(
    path: Path,
    *,
    trait: str | None = None,
    source: str | None = None,
) -> dict[str, float]:
    """ID → numeric value for one trait (and optional source)."""
    out: dict[str, float] = {}
    if not path.exists():
        return out
    df = pd.read_csv(path, sep="\t", dtype=str, comment="#")
    df.columns = [c.strip().lower() for c in df.columns]
    if trait:
        df = df[df["trait"] == trait]
    if source and "source" in df.columns:
        df = df[df["source"] == source]
    for _, r in df.iterrows():
        tid = str(r.get("id") or r.get("sample") or "")
        try:
            out[tid] = float(r["value"])
        except (TypeError, ValueError, KeyError):
            continue
    return out


def build_phenotype_outputs(
    *,
    xlsx: Path,
    euvitis_dict: Path,
    descep_dict: Path,
    info: Path,
    out_dir: Path,
    coverage_dir: Path,
) -> dict[str, Any]:
    panel_ids = load_panel_ids(info)
    dict_df = load_oiv_dictionary(euvitis_dict, descep_dict)
    all_pheno: list[dict] = []
    all_ann_e: list[dict] = []
    all_ann_d: list[dict] = []
    sheet_stats: dict[str, Any] = {}

    xl = pd.ExcelFile(xlsx, engine="openpyxl")
    for sheet in xl.sheet_names:
        if sheet not in SOURCE_SHEETS:
            continue
        ph, ann, st = extract_st1_sheet(xlsx, sheet, panel_ids, dict_df)
        src = SOURCE_SHEETS[sheet]
        sheet_stats[src] = st
        all_pheno.extend(ph)
        if src == "euvitis":
            all_ann_e.extend(ann)
        else:
            all_ann_d.extend(ann)

    pheno_df = pd.DataFrame(all_pheno)
    if len(pheno_df):
        pheno_df = pheno_df.sort_values(["source", "trait", "ID"])
    annot = merge_annot(all_ann_e, all_ann_d)
    annot_df = pd.DataFrame(annot)

    panel_n = {
        "euvitis": sheet_stats.get("euvitis", {}).get("ids_in_panel_n", 0),
        "descep": sheet_stats.get("descep", {}).get("ids_in_panel_n", 0),
    }
    cov = coverage_table(pheno_df, dict_df, panel_n) if len(pheno_df) else pd.DataFrame()

    out_dir.mkdir(parents=True, exist_ok=True)
    coverage_dir.mkdir(parents=True, exist_ok=True)
    pheno_path = out_dir / "phenotype.tsv"
    traits_path = out_dir / "oiv_traits.tsv"
    annot_path = out_dir / "panel" / "sample_annot.tsv"
    annot_path.parent.mkdir(parents=True, exist_ok=True)
    cov_path = coverage_dir / "coverage.tsv"

    cols = ["ID", "trait", "value", "source", "n_obs", "raw"]
    if len(pheno_df):
        pheno_df[cols].to_csv(pheno_path, sep="\t", index=False)
    else:
        pd.DataFrame(columns=cols).to_csv(pheno_path, sep="\t", index=False)
    dict_df.to_csv(traits_path, sep="\t", index=False)
    if len(annot_df):
        annot_df.to_csv(annot_path, sep="\t", index=False)
    else:
        annot_df.to_csv(annot_path, sep="\t", index=False)
    cov.to_csv(cov_path, sep="\t", index=False)

    union_ids = set()
    for st in sheet_stats.values():
        union_ids |= st.get("ids_in_panel") or set()
    n_ge50 = int((cov["n_nonmissing"] >= MIN_PHENO_N).sum()) if len(cov) else 0
    per_src_ge50 = {}
    if len(cov):
        per_src_ge50 = {
            src: int((g["n_nonmissing"] >= MIN_PHENO_N).sum())
            for src, g in cov.groupby("source")
        }

    counts = {
        "xlsx_max_row_euvitis": sheet_stats.get("euvitis", {}).get("xlsx_max_row", 0),
        "xlsx_max_row_descep": sheet_stats.get("descep", {}).get("xlsx_max_row", 0),
        "xlsx_rows_euvitis": sheet_stats.get("euvitis", {}).get("xlsx_rows", 0),
        "xlsx_rows_descep": sheet_stats.get("descep", {}).get("xlsx_rows", 0),
        "uniq_libid_descep": sheet_stats.get("descep", {}).get("uniq_libid", 0),
        "ids_in_panel_euvitis": panel_n["euvitis"],
        "ids_in_panel_descep": panel_n["descep"],
        "ids_in_panel_union": len(union_ids),
        "cells_total": sum(st.get("cells_total", 0) for st in sheet_stats.values()),
        "cells_zero_dropped": sum(st.get("cells_zero_dropped", 0) for st in sheet_stats.values()),
        "cells_multivalue": sum(st.get("cells_multivalue", 0) for st in sheet_stats.values()),
        "cells_out_of_scale_dropped": sum(
            st.get("cells_out_of_scale_dropped", 0) for st in sheet_stats.values()
        ),
        "replicates_aggregated": sum(
            st.get("replicates_aggregated", 0) for st in sheet_stats.values()
        ),
        "traits_total_per_source": (
            cov.groupby("source").size().to_dict() if len(cov) else {}
        ),
        "traits_n_ge_50_per_source": per_src_ge50,
        "traits_n_ge_50_total": n_ge50,
    }
    return {
        "counts": counts,
        "paths": {
            "phenotype": pheno_path,
            "oiv_traits": traits_path,
            "sample_annot": annot_path,
            "coverage": cov_path,
        },
        "n_pheno": len(pheno_df),
    }
